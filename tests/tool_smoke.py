from pathlib import Path
import json, subprocess, sys, tempfile
ROOT=Path(__file__).resolve().parents[1]

def run(*args, cwd=None, ok=(0,)):
    r=subprocess.run([str(x) for x in args],cwd=cwd,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=180)
    if r.returncode not in ok: raise RuntimeError(f"exit={r.returncode} cmd={' '.join(map(str,args))}\n{r.stdout[-5000:]}")
    return r.stdout

with tempfile.TemporaryDirectory(prefix='mms_tool_smoke_') as td:
    td=Path(td)
    print('SMOKE pdf', flush=True)
    # PDF fixture
    import fitz
    pdf=td/'sample.pdf'; d=fitz.open(); p=d.new_page(); p.insert_text((72,72),'Problem statement page 1'); d.save(pdf); d.close()
    run(sys.executable,ROOT/'tools/pdf/scripts/inspect_pdf.py',pdf,'--json')
    txt=td/'pdf.txt'; run(sys.executable,ROOT/'tools/pdf/scripts/extract_text.py',pdf,'--output',txt); assert 'Problem statement' in txt.read_text()
    run(sys.executable,ROOT/'tools/pdf/scripts/check_bounds.py',pdf)
    pdf_render=td/'pdf-render'; run(sys.executable,ROOT/'tools/pdf/scripts/render_pages.py',pdf,'--output-dir',pdf_render); assert (pdf_render/'page-1.png').exists()

    print('SMOKE xlsx', flush=True)
    # XLSX fixture
    import openpyxl
    xlsx=td/'sample.xlsx'; wb=openpyxl.Workbook(); ws=wb.active; ws.append(['x','y']); ws.append([1,2]); ws['C1']='sum'; ws['C2']='=A2+B2'; wb.save(xlsx)
    run(sys.executable,ROOT/'tools/xlsx/scripts/audit_workbook.py',xlsx,'--json')
    run(sys.executable,ROOT/'tools/xlsx/scripts/read_rows.py',xlsx,'--max-row','3')
    recalc=td/'sample-recalc.xlsx'; run(sys.executable,ROOT/'tools/xlsx/scripts/recalc.py',xlsx,'--output',recalc); assert recalc.exists()

    print('SMOKE figure', flush=True)
    # Data/figure fixture
    csv=td/'data.csv'; csv.write_text('x,y,group\n1,2,a\n2,4,a\n3,3,b\n',encoding='utf-8')
    run(sys.executable,ROOT/'tools/figure/scripts/profile_data.py',csv,'--group','group','--json')
    import matplotlib.pyplot as plt
    fig,ax=plt.subplots(figsize=(6,4)); ax.plot([1,2,3],[2,4,3]); ax.set_xlabel('x'); ax.set_ylabel('y')
    sys.path.insert(0,str(ROOT/'tools/figure/scripts')); from export_figure import export_figure
    export_figure(fig,td/'fig',formats=('svg','png'),dpi=200); plt.close(fig)
    run(sys.executable,ROOT/'tools/figure/scripts/check_figure.py',td/'fig.png','--json')
    run(sys.executable,ROOT/'tools/figure/scripts/check_figure.py',td/'fig.svg','--json')
    plot_src=td/'plot.py'; plot_src.write_text("import matplotlib.pyplot as plt\nfig,ax=plt.subplots()\nax.plot([1,2],[2,3])\nax.set_xlabel('x')\nax.set_ylabel('y')\nfig.savefig('out.png')\n",encoding='utf-8')
    run(sys.executable,ROOT/'tools/figure/scripts/validate_source.py',plot_src,'--json')
    gray=td/'fig-gray.png'; run(sys.executable,ROOT/'tools/figure/scripts/visual_qa.py',td/'fig.png','--preview',gray,'--json'); assert gray.exists()

    print('SMOKE docx', flush=True)
    # DOCX fixture + render
    from docx import Document
    docx=td/'sample.docx'; doc=Document(); doc.add_heading('Smoke',1); doc.add_paragraph('Evidence-based result.'); table=doc.add_table(rows=2,cols=2); table.cell(0,0).text='A'; table.cell(0,1).text='B'; doc.save(docx)
    run(sys.executable,ROOT/'tools/docx/scripts/docx_audit.py',docx,'--json')
    run(sys.executable,ROOT/'tools/docx/scripts/inspect_template_format.py',docx,'--json')
    render=td/'render'; out=run(sys.executable,ROOT/'tools/docx/scripts/render_docx.py',docx,'--output-dir',render); assert (render/'page-1.png').exists()
    run(sys.executable,ROOT/'tools/docx/scripts/self_check.py',docx)

    print('SMOKE latex', flush=True)
    # LaTeX CJK doctor/build/bind/validate
    tex=td/'paper/main.tex'; tex.parent.mkdir(); run(sys.executable,ROOT/'tools/latex/scripts/latex_paper.py','init-cjk',tex)
    run(sys.executable,ROOT/'tools/latex/scripts/latex_paper.py','doctor','--engine','xelatex')
    run(sys.executable,ROOT/'tools/latex/scripts/latex_paper.py','build',tex,'--engine','xelatex')
    run(sys.executable,ROOT/'tools/latex/scripts/latex_paper.py','validate',tex)
    manifest=td/'paper/latex-project.json'; run(sys.executable,ROOT/'tools/latex/scripts/latex_paper.py','bind',tex.parent,'--output',manifest); assert manifest.exists()
    init_dir=td/'latex-init'; run(sys.executable,ROOT/'tools/latex/scripts/latex_paper.py','init',init_dir,'--contest','generic'); assert (init_dir/'main.tex').exists() and (init_dir/'references.bib').exists()

    print('SMOKE reproducibility', flush=True)
    manifest=td/'run-manifest.json'
    run(sys.executable,ROOT/'tools/reproducibility/scripts/run_manifest.py','doctor','--features','data')
    run(sys.executable,ROOT/'tools/reproducibility/scripts/run_manifest.py','create','--output',manifest,'--run-id','smoke','--command','python smoke.py','--seed','42','--input',csv,'--artifact',td/'fig.png','--package','matplotlib')
    run(sys.executable,ROOT/'tools/reproducibility/scripts/run_manifest.py','verify',manifest)

    print('SMOKE search', flush=True)
    # Network-independent search logic test
    assert 'SELF_TEST_PASS' in run(sys.executable,ROOT/'tools/paper-search/scripts/hybrid_scholar.py','--self-test')

print('TOOL_SMOKE_PASS')
