#!/usr/bin/env python3
import argparse, csv, json
from pathlib import Path
import openpyxl

def main():
    ap=argparse.ArgumentParser(description='Read a bounded XLSX range for factual inspection.')
    ap.add_argument('xlsx'); ap.add_argument('--sheet'); ap.add_argument('--min-row',type=int,default=1); ap.add_argument('--max-row',type=int,default=20)
    ap.add_argument('--min-col',type=int,default=1); ap.add_argument('--max-col',type=int,default=20); ap.add_argument('--json',action='store_true')
    a=ap.parse_args(); wb=openpyxl.load_workbook(a.xlsx,data_only=False,read_only=True); ws=wb[a.sheet] if a.sheet else wb[wb.sheetnames[0]]
    rows=[[c.value for c in row] for row in ws.iter_rows(min_row=a.min_row,max_row=min(a.max_row,ws.max_row),min_col=a.min_col,max_col=min(a.max_col,ws.max_column))]
    if a.json: print(json.dumps({'sheet':ws.title,'rows':rows},ensure_ascii=False,indent=2,default=str))
    else:
        w=csv.writer(__import__('sys').stdout); w.writerows(rows)
if __name__=='__main__': main()
