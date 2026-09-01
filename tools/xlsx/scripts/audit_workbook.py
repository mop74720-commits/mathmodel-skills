#!/usr/bin/env python3
import argparse, json, hashlib
from pathlib import Path
from collections import Counter
import openpyxl

def sha256(path):
    h=hashlib.sha256();
    with open(path,'rb') as f:
        for c in iter(lambda:f.read(1024*1024),b''): h.update(c)
    return h.hexdigest()

def inspect_sheet(ws, max_scan):
    nonempty=[]; formulas=0; errors=[]
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                nonempty.append(c.value)
                if c.data_type=='f': formulas+=1
                if c.data_type=='e': errors.append(c.coordinate)
        if len(nonempty)>=max_scan: break
    headers=[]
    if ws.max_row:
        headers=[ws.cell(1,j).value for j in range(1,ws.max_column+1)]
    return {
      'title':ws.title,'rows':ws.max_row,'cols':ws.max_column,'headers':headers,
      'formula_cells':formulas,'error_cells':errors[:50], 'merged_ranges':[str(x) for x in ws.merged_cells.ranges],
      'hidden':ws.sheet_state!='visible','freeze_panes':str(ws.freeze_panes) if ws.freeze_panes else None,
    }

def main():
    ap=argparse.ArgumentParser(description='Audit XLSX structure and formula/error cells without changing the workbook.')
    ap.add_argument('xlsx'); ap.add_argument('--json',action='store_true'); ap.add_argument('--max-scan',type=int,default=100000)
    a=ap.parse_args(); p=Path(a.xlsx)
    wb=openpyxl.load_workbook(p,data_only=False,read_only=False)
    out={'path':str(p),'sha256':sha256(p),'sheets':[inspect_sheet(ws,a.max_scan) for ws in wb.worksheets],
         'defined_names':sorted(str(x) for x in wb.defined_names)}
    if a.json: print(json.dumps(out,ensure_ascii=False,indent=2,default=str))
    else:
        print(f"Workbook {p.name}: {len(out['sheets'])} sheets sha256={out['sha256']}")
        for s in out['sheets']: print(f"{s['title']}: {s['rows']}x{s['cols']} formulas={s['formula_cells']} errors={len(s['error_cells'])}")
if __name__=='__main__': main()
